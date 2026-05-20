import docx
from docx.oxml.ns import qn
from docx.table import _Cell

def _format_docx_table(tbl):
    grid = {}
    for r_idx, tr in enumerate(tbl._tbl.tr_lst):
        c_idx = 0
        for tc in tr.tc_lst:
            while (r_idx, c_idx) in grid:
                c_idx += 1
            
            tcPr = tc.tcPr
            grid_span = 1
            v_merge = None
            if tcPr is not None:
                gridSpan_elem = tcPr.find(qn('w:gridSpan'))
                if gridSpan_elem is not None:
                    val = gridSpan_elem.get(qn('w:val'))
                    if val:
                        grid_span = int(val)
                
                vMerge_elem = tcPr.find(qn('w:vMerge'))
                if vMerge_elem is not None:
                    val = vMerge_elem.get(qn('w:val'))
                    v_merge = val if val else 'continue'
                    
            if v_merge == 'continue':
                text = ""
            else:
                text = _Cell(tc, tbl).text.replace("\n", " ").replace("\r", " ").strip()
                
            for i in range(grid_span):
                grid[(r_idx, c_idx + i)] = text if i == 0 else ""
                
            c_idx += grid_span
            
    if not grid:
        return ""
        
    max_r = max((r for r, c in grid.keys()), default=-1)
    max_c = max((c for r, c in grid.keys()), default=-1)
    
    if max_r < 0 or max_c < 0:
        return ""
    
    all_rows = []
    for r in range(max_r + 1):
        row_data = []
        for c in range(max_c + 1):
            row_data.append(grid.get((r, c), ""))
        all_rows.append(row_data)
        
    n_cols = max_c + 1
    md_lines = []
    for ri, row in enumerate(all_rows):
        md_lines.append("| " + " | ".join(row) + " |")
        if ri == 0:
            md_lines.append("| " + " | ".join(["---"] * n_cols) + " |")
    return "\n".join(md_lines)

doc = docx.Document()
table = doc.add_table(rows=2, cols=2)
a = table.cell(0, 0)
b = table.cell(0, 1)
a.merge(b)
a.text = "Merged Header"
table.cell(1, 0).text = "Cell 1"
table.cell(1, 1).text = "Cell 2"

print(_format_docx_table(table))
